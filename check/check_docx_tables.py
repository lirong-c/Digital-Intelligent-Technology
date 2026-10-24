# -*- coding: utf-8 -*-
"""校验：稿件 docx 内的表格是否与 expected/tables.json、expected/vif.csv 一致（无需原始数据）。"""
import json, os, sys
HERE = os.path.dirname(os.path.abspath(__file__)); ROOT = os.path.dirname(HERE)
DOCX = os.path.join(ROOT, "expected", "manuscript.docx")
EXP = json.load(open(os.path.join(ROOT, "expected", "tables.json"), encoding="utf-8"))
VIF = [r for r in __import__("csv").reader(open(os.path.join(ROOT, "expected", "vif.csv"), encoding="utf-8-sig"))]

try:
    from docx import Document
except ImportError:
    sys.exit("需要 python-docx：pip install python-docx")

doc = Document(DOCX)
def norm(x):
    """-0.0000 与 0.0000 视为相同（稿件中将负零统一显示为正零）。"""
    x = str(x).strip()
    return "0.0000" if x in ("-0.0000", "-0", "-0.000") else x
def cells(t): return [[norm(c.text) for c in r.cells] for r in t.rows]
YEARS = (2007, 2010, 2013, 2016, 2019, 2020, 2022, 2024)
tbl5 = [EXP["表8"][0]] + [r for r in EXP["表8"][1:] if int(float(r[0])) in YEARS]
CHECKS = [
    (3,  "表4  描述性统计", [[str(c) for c in r] for r in EXP["表7"]],  None),
    (4,  "表5  年度演变",   [[str(c) for c in r] for r in tbl5],         None),
    (5,  "表6  相关系数",   [[str(c) for c in r] for r in EXP["表9"]],   None),
    (7,  "表8  基准回归",   [[str(c) for c in r] for r in EXP["表11"]],  None),
    (12, "表13 稳健性",     [[str(c) for c in r] for r in EXP["表12"]],  None),
    (13, "表14 中介",       [[str(c) for c in r] for r in EXP["表13"]],  None),
    (14, "表15 空间权重",   [[str(c) for c in r] for r in EXP["表14"]],  None),
]
DERIVED = [(17, "表18 行业异质性", 1, 3), (18, "表19 区域异质性", 3, 6)]
total = same = 0
for idx, name, ref, _ in CHECKS:
    got = cells(doc.tables[idx])
    if len(got) != len(ref) or any(len(a) != len(b) for a, b in zip(got, ref)):
        print("✗ %s 形状不符" % name); continue
    bad = [(i, j, a, b) for i, (ra, rb) in enumerate(zip(got, ref))
           for j, (a, b) in enumerate(zip(ra, rb)) if norm(a) != norm(b)]
    total += sum(len(r) for r in ref); same += sum(len(r) for r in ref) - len(bad)
    print(("✓ " if not bad else "✗ ") + "%-12s 差异 %d" % (name, len(bad)))
    for b in bad[:4]: print("     行%d列%d docx=%r 基准=%r" % b)
base = EXP["表15_表16"]
for idx, name, lo, hi in DERIVED:
    got = cells(doc.tables[idx])
    want = [["子样本", "系数", "t值", "N", "组间差异检验"]]
    want += [[r[1], r[2], r[3], str(r[4]), r[5] if len(r) > 5 and r[5] else "—"] for r in base[lo:hi]]
    ok = got == want
    total += sum(len(r) for r in want); same += sum(len(r) for r in want) if ok else 0
    print(("✓ " if ok else "✗ ") + name)
# VIF
got = cells(doc.tables[6]); want = [[str(c) for c in r] for r in VIF]
diff = [(a, b) for ra, rb in zip(got, want) for a, b in zip(ra, rb) if a != b]
print(("✓ " if not diff else "△ ") + "表7  VIF（差异 %d 处，说明见 README 第 7 节）" % len(diff))
print("\n合计: %d/%d 单元格一致" % (same, total))
